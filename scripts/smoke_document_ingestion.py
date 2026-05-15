"""Smoke tests for Phase AI-9 universal document ingestion.

The fixtures are generated in a temporary folder so this script does not touch
training bills, Output/Template.xlsx, Dropbox, or any AI provider.
"""

from __future__ import annotations

import sys
import tempfile
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from webapp.backend import settings  # noqa: E402
from webapp.backend.services import document_ingestion  # noqa: E402


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _write_simple_pdf(path: Path, text: str) -> None:
    content = f"BT /F1 12 Tf 72 720 Td ({_pdf_escape(text)}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(data))
        data.extend(f"{idx} 0 obj\n".encode("ascii"))
        data.extend(obj)
        data.extend(b"\nendobj\n")
    xref_offset = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    path.write_bytes(bytes(data))


def _write_docx(path: Path) -> None:
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>Word invoice fixture</w:t></w:r></w:p>
    <w:tbl>
      <w:tr><w:tc><w:p><w:r><w:t>Item</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Amount</w:t></w:r></w:p></w:tc></w:tr>
      <w:tr><w:tc><w:p><w:r><w:t>Repair</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>42.00</w:t></w:r></w:p></w:tc></w:tr>
    </w:tbl>
  </w:body>
</w:document>"""
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types />")
        zf.writestr("word/document.xml", document_xml)


def _write_image(path: Path) -> None:
    from PIL import Image, ImageDraw  # type: ignore

    img = Image.new("RGB", (420, 180), "white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 40), "Screenshot invoice fixture\nTotal 18.90", fill="black")
    img.save(path)


def _write_xlsx(path: Path) -> None:
    import openpyxl  # type: ignore

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.append(["Invoice Number", "Total"])
    ws.append(["XL-100", 55.25])
    wb.save(path)


def main() -> int:
    template_path = settings.RESMAN_TEMPLATE
    template_mtime = template_path.stat().st_mtime if template_path.exists() else None
    with tempfile.TemporaryDirectory(prefix="document_ingestion_") as tmp:
        base = Path(tmp)
        pdf = base / "digital_invoice.pdf"
        _write_simple_pdf(pdf, "Invoice Number PDF-123 Account 555 Total 12.34 " * 5)
        pdf_candidate = document_ingestion.ingest_document(pdf)
        _assert(pdf_candidate.source_type == "pdf_digital", "PDF should be classified as pdf_digital")
        _assert("PDF-123" in pdf_candidate.document_text, "PDF text should be extracted")
        _assert(pdf_candidate.pages, "PDF should include page candidates")
        _assert(pdf_candidate.page_count == 1, "PDF should expose page_count")
        _assert(pdf_candidate.text_quality_score > 0, "PDF should expose text_quality_score")
        _assert(pdf_candidate.extraction_quality.get("label") in {"high", "medium", "low"}, "PDF quality label should be present")
        rehydrated = document_ingestion.document_candidate_from_dict(pdf_candidate.to_dict())
        _assert(rehydrated.source_file == pdf_candidate.source_file, "DocumentCandidate should round-trip through dict")

        csv_path = base / "invoice.csv"
        csv_path.write_text("invoice,total\nCSV-1,88.10\n", encoding="utf-8")
        csv_candidate = document_ingestion.ingest_document(csv_path)
        _assert(csv_candidate.source_type == "csv", "CSV should be classified as csv")
        _assert(csv_candidate.tables and csv_candidate.tables[0].columns, "CSV table should be extracted")

        xlsx_path = base / "invoice.xlsx"
        _write_xlsx(xlsx_path)
        xlsx_candidate = document_ingestion.ingest_document(xlsx_path)
        _assert(xlsx_candidate.source_type == "excel", "XLSX should be classified as excel")
        _assert(xlsx_candidate.tables, "XLSX tables should be extracted")
        _assert("XL-100" in xlsx_candidate.document_text, "XLSX text summary should include cells")

        docx_path = base / "invoice.docx"
        _write_docx(docx_path)
        docx_candidate = document_ingestion.ingest_document(docx_path)
        _assert(docx_candidate.source_type == "word", "DOCX should be classified as word")
        _assert("Word invoice fixture" in docx_candidate.document_text, "DOCX text should be extracted")
        _assert(docx_candidate.tables, "DOCX tables should be extracted")

        image_path = base / "screenshot_invoice.png"
        _write_image(image_path)
        image_candidate = document_ingestion.ingest_document(image_path)
        _assert(image_candidate.source_type == "screenshot", "Screenshot PNG should be classified as screenshot")
        _assert(image_candidate.pages and image_candidate.images, "Image should create one page and image reference")
        _assert(image_candidate.page_count == 1, "Image should expose one page")

        unsupported_path = base / "invoice.bin"
        unsupported_path.write_bytes(b"\x00\x01\x02")
        unsupported = document_ingestion.ingest_document(unsupported_path)
        _assert(unsupported.source_type == "unknown", "Unknown extension should be source_type unknown")
        _assert("unsupported_file_type" in unsupported.warnings, "Unsupported file should include warning")

        if template_path.exists():
            protected = document_ingestion.ingest_document(template_path)
            _assert(protected.source_type == "internal_template", "Output/Template.xlsx should be protected")
            _assert(
                "internal_resman_template_not_ingested" in protected.warnings,
                "Internal template should return a friendly ignored warning",
            )

    if template_mtime is not None:
        _assert(template_path.stat().st_mtime == template_mtime, "Output/Template.xlsx must not be modified")

    print("document ingestion smoke: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
